"""Tests for hippo.ingest.readers: every supported format, plus what gets skipped."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from hippo.ingest import readers
from hippo.ingest.html_text import html_to_text
from hippo.ingest.readers import Document, is_probably_binary, is_supported, read_file, read_zip

# ------------------------------------------------------------------ helpers


def minimal_pdf(text: str) -> bytes:
    """A one-page PDF with `text` drawn in Helvetica; pypdf can extract it back."""
    stream = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def minimal_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("The company", level=2)
    document.add_paragraph("Acme Robotics is headquartered in Boulder.")
    document.add_paragraph("Acme Robotics builds robot arms.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Orion arm"
    table.rows[0].cells[1].text = "12 kilograms"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def minimal_epub() -> bytes:
    """Two chapters whose spine order (b then a) differs from alphabetical order."""
    opf = """<?xml version="1.0"?>
    <package xmlns="http://www.idpf.org/2007/opf" version="2.0">
      <manifest>
        <item id="a" href="a.xhtml" media-type="application/xhtml+xml"/>
        <item id="b" href="b.xhtml" media-type="application/xhtml+xml"/>
      </manifest>
      <spine><itemref idref="b"/><itemref idref="a"/></spine>
    </package>"""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr("OEBPS/content.opf", opf)
        book.writestr("OEBPS/a.xhtml", "<html><body><h1>Second</h1><p>Text of chapter two.</p></body></html>")
        book.writestr("OEBPS/b.xhtml", "<html><body><h1>First</h1><p>Text of chapter one.</p></body></html>")
    return buffer.getvalue()


# ------------------------------------------------------------ plain files


def test_markdown_file_is_prose(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("# Title\n\nAcme Robotics is headquartered in Boulder.\n")
    docs = read_file(path)
    assert len(docs) == 1
    doc = docs[0]
    assert isinstance(doc, Document)
    assert doc.title == "notes.md"
    assert doc.is_code is False
    assert "Boulder" in doc.text
    assert doc.path == str(path)


def test_python_file_is_code(tmp_path: Path) -> None:
    path = tmp_path / "tool.py"
    path.write_text("def hello():\n    return 1\n")
    (doc,) = read_file(path)
    assert doc.is_code is True
    assert doc.title == "tool.py"


def test_empty_file_gives_no_document(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n")
    assert read_file(path) == []


def test_utf8_bom_is_removed(tmp_path: Path) -> None:
    path = tmp_path / "bom.txt"
    path.write_bytes(b"\xef\xbb\xbfHello there.")
    (doc,) = read_file(path)
    assert doc.text == "Hello there."


def test_extensionless_text_is_supported(tmp_path: Path) -> None:
    readme = tmp_path / "README"
    readme.write_text("This project builds robot arms.")
    assert is_supported(readme)
    (doc,) = read_file(readme)
    assert doc.title == "README"
    assert doc.is_code is False


def test_extensionless_binary_is_not_supported(tmp_path: Path) -> None:
    blob = tmp_path / "blob"
    blob.write_bytes(b"\x00\x01\x02\xff\xfe" * 100)
    assert not is_supported(blob)
    assert read_file(blob) == []


def test_unknown_extension_is_not_supported(tmp_path: Path) -> None:
    image = tmp_path / "photo.png"
    image.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 50)
    assert not is_supported(image)
    assert is_supported(tmp_path / "anything.pdf")  # decided by name; the file need not exist


def test_binary_with_known_text_extension_is_skipped(tmp_path: Path) -> None:
    path = tmp_path / "weird.txt"
    path.write_bytes(b"abc\x00def")
    assert read_file(path) == []


def test_is_probably_binary() -> None:
    assert is_probably_binary(b"\x00\x00\x00")
    assert is_probably_binary(b"caf\xe9 latin-1 bytes")  # not UTF-8
    assert not is_probably_binary("plain text, with ünïcödé".encode())
    assert not is_probably_binary(b"")


# ------------------------------------------------------------ rich files


def test_html_strips_tags_and_keeps_headings() -> None:
    html = """<html><head><title>x</title><style>p{}</style></head><body>
    <script>alert(1)</script>
    <h2>The   company</h2>
    <p>Acme Robotics is <b>headquartered</b> in Boulder.</p>
    <p>Second paragraph.</p>
    </body></html>"""
    text = html_to_text(html)
    assert "alert" not in text and "p{}" not in text and "<" not in text
    assert "## The company" in text
    assert "Acme Robotics is headquartered in Boulder." in text
    assert "\n\nSecond paragraph." in text


def test_html_file(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text("<h1>Hi</h1><p>Body text.</p>")
    (doc,) = read_file(path)
    assert doc.text == "# Hi\n\nBody text."
    assert doc.is_code is False


def test_pdf_file(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    path.write_bytes(minimal_pdf("Acme Robotics is headquartered in Boulder."))
    (doc,) = read_file(path)
    assert doc.title == "guide.pdf"
    assert "Acme Robotics is headquartered in Boulder." in doc.text


def test_broken_pdf_raises_read_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(readers.ReadError):
        read_file(path)


def test_docx_file(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    path.write_bytes(minimal_docx())
    (doc,) = read_file(path)
    assert "## The company" in doc.text
    assert "Acme Robotics is headquartered in Boulder." in doc.text
    assert "Orion arm | 12 kilograms" in doc.text


def test_epub_follows_spine_order(tmp_path: Path) -> None:
    path = tmp_path / "book.epub"
    path.write_bytes(minimal_epub())
    (doc,) = read_file(path)
    assert doc.text.index("# First") < doc.text.index("# Second")
    assert "Text of chapter one." in doc.text


# ------------------------------------------------------------------- zips


def build_zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def test_read_zip_reads_supported_members_and_skips_junk(tmp_path: Path) -> None:
    path = tmp_path / "bundle.zip"
    path.write_bytes(
        build_zip(
            {
                "docs/guide.md": b"# Guide\n\nAcme Robotics builds robot arms.",
                "src/app.py": b"print('hi')\n",
                "README": b"Extensionless readme text.",
                "node_modules/x/index.js": b"junk",
                ".git/config": b"junk",
                "__MACOSX/._guide.md": b"junk",
                "images/logo.png": b"\x89PNG\x00\x00",
                "data/blob.txt": b"text\x00with nul",
                "empty.txt": b"",
            }
        )
    )
    docs = read_zip(path, "bundle.zip")
    titles = [d.title for d in docs]
    assert titles == ["README", "docs/guide.md", "src/app.py"]
    by_title = {d.title: d for d in docs}
    assert by_title["src/app.py"].is_code is True
    assert by_title["docs/guide.md"].is_code is False
    assert by_title["docs/guide.md"].path == "bundle.zip:docs/guide.md"


def test_read_zip_skips_files_over_the_size_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(readers, "MAX_FILE_BYTES", 100)
    path = tmp_path / "bundle.zip"
    path.write_bytes(build_zip({"small.txt": b"small", "big.txt": b"x" * 200}))
    assert [d.title for d in read_zip(path, "bundle.zip")] == ["small.txt"]


def test_read_zip_keeps_going_when_one_member_is_broken(tmp_path: Path) -> None:
    path = tmp_path / "bundle.zip"
    path.write_bytes(build_zip({"bad.pdf": b"not a pdf", "good.txt": b"Fine text."}))
    assert [d.title for d in read_zip(path, "bundle.zip")] == ["good.txt"]


# ---------------------------------------------------------------- budgets
# A small file can unpack into a huge amount of text (epub, zip, pdf streams). Every reader
# counts what it produces against a TextBudget and stops with TooLarge, which is a ReadError
# that read_zip does *not* swallow as "one broken member".


def test_text_budget_raises_too_large_once_the_total_passes_the_limit() -> None:
    budget = readers.TextBudget(limit=10)
    budget.add(6, "a.txt")
    budget.add(4, "b.txt")  # exactly at the limit is fine
    with pytest.raises(readers.TooLarge, match="too large"):
        budget.add(1, "c.txt")
    assert issubclass(readers.TooLarge, readers.ReadError)


def test_read_bytes_counts_plain_text_against_the_budget() -> None:
    budget = readers.TextBudget(limit=5)
    with pytest.raises(readers.TooLarge, match="notes.md"):
        readers.read_bytes(b"more than five", "notes.md", budget=budget)
    assert readers.read_bytes(b"tiny", "notes.md", budget=readers.TextBudget(limit=5))


def test_read_zip_shares_one_budget_across_members_and_does_not_swallow_it(tmp_path: Path) -> None:
    path = tmp_path / "bundle.zip"
    path.write_bytes(build_zip({"a.txt": b"x" * 40, "b.txt": b"y" * 40}))
    assert len(read_zip(path, "bundle.zip", readers.TextBudget(limit=80))) == 2
    with pytest.raises(readers.TooLarge):
        read_zip(path, "bundle.zip", readers.TextBudget(limit=79))


def test_read_zip_refuses_too_many_members_or_too_much_unpacked_data(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "bundle.zip"
    path.write_bytes(
        build_zip({"a.txt": b"x" * 30, "b.txt": b"y" * 30, "c.txt": b"z" * 30, "junk.png": b"\x00"})
    )
    monkeypatch.setattr(readers, "MAX_ZIP_MEMBERS", 2)
    with pytest.raises(readers.TooLarge, match="3 readable files"):
        read_zip(path, "bundle.zip")
    monkeypatch.setattr(readers, "MAX_ZIP_MEMBERS", 5_000)
    monkeypatch.setattr(readers, "MAX_ZIP_TOTAL_BYTES", 89)
    with pytest.raises(readers.TooLarge, match="unpack to 90 bytes"):
        read_zip(path, "bundle.zip")
    monkeypatch.setattr(readers, "MAX_ZIP_TOTAL_BYTES", 90)
    assert len(read_zip(path, "bundle.zip")) == 3


def test_epub_chapters_count_one_by_one_and_oversized_chapters_are_skipped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "book.epub"
    path.write_bytes(minimal_epub())
    # The first chapter in spine order (b.xhtml) already breaks a tiny budget.
    with pytest.raises(readers.TooLarge, match="b.xhtml"):
        read_file(path, readers.TextBudget(limit=10))
    # Chapters bigger than MAX_FILE_BYTES are skipped without being read: a.xhtml ("Second") is one byte
    # longer than b.xhtml ("First"), so a cap of exactly b's size keeps chapter one and drops chapter two.
    with zipfile.ZipFile(io.BytesIO(minimal_epub())) as book:
        monkeypatch.setattr(readers, "MAX_FILE_BYTES", book.getinfo("OEBPS/b.xhtml").file_size)
    (doc,) = read_file(path)
    assert "chapter one" in doc.text and "chapter two" not in doc.text


def test_pdf_pages_count_against_the_budget(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    path.write_bytes(minimal_pdf("Acme Robotics is headquartered in Boulder."))
    with pytest.raises(readers.TooLarge):
        read_file(path, readers.TextBudget(limit=10))


def test_docx_with_a_huge_document_body_is_refused_before_parsing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "guide.docx"
    path.write_bytes(minimal_docx())
    monkeypatch.setattr(readers, "MAX_DOCX_XML_BYTES", 100)
    with pytest.raises(readers.TooLarge, match="document body"):
        read_file(path)


def test_walk_repo_does_not_swallow_too_large(tmp_path: Path) -> None:
    from hippo.ingest.repos import walk_repo

    (tmp_path / "a.md").write_text("x" * 40)
    (tmp_path / "b.md").write_text("y" * 40)
    assert len(walk_repo(tmp_path, readers.TextBudget(limit=80))) == 2
    with pytest.raises(readers.TooLarge):
        walk_repo(tmp_path, readers.TextBudget(limit=79))
