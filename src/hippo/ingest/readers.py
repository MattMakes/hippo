"""
Turning files into plain text.

Whatever you give hippo (a PDF, a Word file, an e-book, a web page, a source
file, a zip of any of those) ends up as one or more `Document`s: a title, the
text, and a flag saying whether the text is code (code is chunked by lines,
prose by paragraphs; see chunker.py).

Rules of thumb used here:

* Rich formats (.pdf .docx .epub .html) get a small dedicated reader.
* Anything with a known text/code extension is read as UTF-8.
* A file with *no* extension is accepted if it decodes as UTF-8 and has no
  NUL bytes (a README, a Makefile, a LICENSE...). Everything else is skipped.
* Inside zips we ignore the usual junk folders (.git, node_modules, ...),
  hidden folders, and anything bigger than MAX_FILE_BYTES.
* Every reader counts the text it produces against a `TextBudget` and stops
  with `TooLarge` past MAX_TEXT_CHARS. A small file can unpack into gigabytes
  of text (an epub or zip "bomb"), and every 1500 characters cost two LLM
  calls later, so the cap is on the text that comes out, not the bytes that
  go in. Zips also have a member-count and a total-size budget.
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from .html_text import html_to_text

log = logging.getLogger(__name__)

IGNORED_DIRS = {
    ".git",
    "node_modules",
    "__pycache__",
    ".venv",
    "venv",
    "dist",
    "build",
    "target",
    ".idea",
    ".vscode",
    "vendor",
}
MAX_FILE_BYTES = 2_000_000
# Text budget per source (all files together). 20 M chars is ~13,000 passages, about the most
# a local model can index in a day; pipeline.py passes the configured value (HIPPO_MAX_TEXT_CHARS).
MAX_TEXT_CHARS = 20_000_000
MAX_ZIP_MEMBERS = 5_000  # readable members per archive
MAX_ZIP_TOTAL_BYTES = 50_000_000  # unpacked size of the readable members together
MAX_DOCX_XML_BYTES = 20_000_000  # size of word/document.xml; python-docx parses it whole

# Prose: chunked by headings/paragraphs/sentences.
PROSE_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".text"}
# Rich formats with their own reader.
RICH_EXTENSIONS = {".pdf", ".docx", ".epub", ".html", ".htm"}
# Code and config: chunked by lines. Generous on purpose; the cost of a wrong
# guess is only a slightly odd chunk boundary.
CODE_EXTENSIONS = {
    ".py", ".pyi", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".kt", ".kts",
    ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".cs", ".rb", ".php", ".swift", ".scala", ".m",
    ".sh", ".bash", ".zsh", ".fish", ".ps1", ".bat", ".sql", ".r", ".jl", ".lua", ".pl", ".pm", ".ex",
    ".exs", ".erl", ".hs", ".ml", ".clj", ".dart", ".groovy", ".gradle", ".cmake", ".mk", ".makefile",
    ".yaml", ".yml", ".toml", ".json", ".jsonl", ".xml", ".css", ".scss", ".sass", ".less", ".ini",
    ".cfg", ".conf", ".properties", ".proto", ".graphql", ".tf", ".hcl", ".vue", ".svelte",
    ".csv", ".tsv", ".dockerfile", ".lock",
}  # fmt: skip
# The subsets of CODE_EXTENSIONS the code graph has a grammar (or a parser) for; see `lang_of`.
PYTHON_EXTENSIONS = {".py", ".pyi"}
TYPESCRIPT_EXTENSIONS = {".ts", ".tsx", ".js", ".jsx", ".mjs", ".cjs"}
GO_EXTENSIONS = {".go"}
CSHARP_EXTENSIONS = {".cs"}
RUST_EXTENSIONS = {".rs"}
SQL_EXTENSIONS = {".sql"}
# Extensionless files that we know are text.
KNOWN_TEXT_NAMES = {"makefile", "dockerfile", "license", "readme", "notice", "authors", "changelog"}
# Files we know by their *whole* name, extension and all. `go.mod` is here rather than in
# CODE_EXTENSIONS because `.mod` belongs to Fortran and half a dozen other things too; the
# Go walker reads its `module` line (through `LanguageRules.source_setup`) to turn an import
# path into a directory, and it can only do that if the file is read at all.
KNOWN_TEXT_FILENAMES = {"go.mod"}

BINARY_SNIFF_BYTES = 8_000


@dataclass
class Document:
    """One readable thing: a file, a zip member, a pasted text."""

    title: str
    text: str
    path: str
    is_code: bool


class ReadError(ValueError):
    """The file exists but we could not get text out of it."""


class TooLarge(ReadError):
    """The source holds more text than we are willing to index. Never swallowed: it fails the whole source."""


@dataclass
class TextBudget:
    """
    A running count of the characters a source has produced so far.
    `add` raises TooLarge once the count passes `limit`. One budget is shared
    by every file of a source, so the total is what is capped.
    """

    limit: int = MAX_TEXT_CHARS
    used: int = 0

    def add(self, chars: int, where: str) -> None:
        self.used += chars
        if self.used > self.limit:
            raise TooLarge(
                f"too large: more than {self.limit:,} characters of text (reached at {where}). "
                "Split the source into smaller parts, or raise HIPPO_MAX_TEXT_CHARS."
            )


# ------------------------------------------------------------- questions


def is_probably_binary(data: bytes) -> bool:
    """True when the start of `data` has NUL bytes or is not valid UTF-8."""
    head = data[:BINARY_SNIFF_BYTES]
    if b"\x00" in head:
        return True
    try:
        # A cut in the middle of a multi-byte character would look invalid, so
        # ignore a partial tail: decode with a tolerant fallback on the last bytes.
        head.decode("utf-8")
    except UnicodeDecodeError as err:
        return err.start < len(head) - 4
    return False


def is_code_name(name: str) -> bool:
    return _suffix(name) in CODE_EXTENSIONS


def lang_of(name: str) -> str | None:
    """
    Which language the code graph knows this file as -- one of the five, or "sql" -- and
    None for every other file (Ruby, C, prose, anything with no suffix).

    Naming a language here is not the same as parsing it: a language may be registered
    (suffix, grammar, comment style) before its walker exists, and `codegraph.extract` then
    keeps today's line windows for its files. Keyed on the same suffixes `is_code_name` uses,
    so a file can never be code for the chunker and unknown to the extractor.
    `.js/.jsx/.mjs/.cjs` are "typescript": the TSX grammar parses plain JavaScript cleanly
    (R4 T9), so there is no third grammar. `codegraph.model.lang_of` is the same table on the
    other side of the dependency line; `test_ingest_readers.py` pins the two together.
    """
    suffix = _suffix(name)
    if suffix not in CODE_EXTENSIONS:
        return None
    if suffix in PYTHON_EXTENSIONS:
        return "python"
    if suffix in TYPESCRIPT_EXTENSIONS:
        return "typescript"
    if suffix in GO_EXTENSIONS:
        return "go"
    if suffix in CSHARP_EXTENSIONS:
        return "csharp"
    if suffix in RUST_EXTENSIONS:
        return "rust"
    return "sql" if suffix in SQL_EXTENSIONS else None


def is_supported_name(name: str) -> bool:
    """Decide by name alone (used for zip members, uploads and repo files)."""
    suffix = _suffix(name)
    if suffix in PROSE_EXTENSIONS or suffix in RICH_EXTENSIONS or suffix in CODE_EXTENSIONS:
        return True
    stem = PurePosixPath(name).name.lower()
    if stem in KNOWN_TEXT_FILENAMES:
        return True
    return suffix == "" and stem in KNOWN_TEXT_NAMES


def is_plain_prose_name(name: str) -> bool:
    """True for a stored file name a managed plain-prose build can read.

    Closed on `PROSE_EXTENSIONS` and decided by name alone: an extensionless or
    sniffed-as-text file is readable but has no honest plain-prose original, and
    a claimed content type is never consulted.
    """
    return _suffix(str(name).replace("\\", "/")) in PROSE_EXTENSIONS


def is_supported(path: Path) -> bool:
    """True if `read_file` would try to read it. Extensionless files are sniffed for text."""
    if is_supported_name(path.name):
        return True
    if path.suffix or not path.is_file():
        return False
    try:
        with path.open("rb") as handle:
            return not is_probably_binary(handle.read(BINARY_SNIFF_BYTES))
    except OSError:
        return False


# --------------------------------------------------------------- reading


def read_file(path: Path, budget: TextBudget | None = None) -> list[Document]:
    """Read one file from disk. The document title is the file name."""
    return read_path(path, path.name, budget)


def read_path(path: Path, title: str, budget: TextBudget | None = None) -> list[Document]:
    """Like read_file, but you choose the title (repos use the path relative to the root)."""
    data = path.read_bytes()
    return read_bytes(data, title, path=str(path), budget=budget)


def read_bytes(
    data: bytes, name: str, path: str | None = None, budget: TextBudget | None = None
) -> list[Document]:
    """
    Read from memory. `name` decides the format (by extension) and becomes the title.
    `budget` is shared across the files of one source; without one, this file gets a fresh budget.
    """
    budget = budget or TextBudget()
    suffix = _suffix(name)
    where = path or name
    if suffix == ".pdf":
        text = _read_pdf(data, budget, where)
    elif suffix == ".docx":
        text = _read_docx(data)
    elif suffix == ".epub":
        text = _read_epub(data, budget, where)
    elif suffix in (".html", ".htm"):
        text = html_to_text(_decode(data, strict=False))
    elif suffix in PROSE_EXTENSIONS or suffix in CODE_EXTENSIONS:
        if is_probably_binary(data):
            log.info("Skipping %s: looks binary", where)
            return []
        text = _decode(data, strict=False)
    else:
        # No known extension: accept only clean UTF-8 text.
        if is_probably_binary(data):
            return []
        text = _decode(data, strict=True)
    text = text.strip()
    if not text:
        return []
    if suffix not in (".pdf", ".epub"):  # those two counted page by page / chapter by chapter
        budget.add(len(text), where)
    return [Document(title=name, text=text, path=where, is_code=suffix in CODE_EXTENSIONS)]


def read_zip(path: Path, name: str, budget: TextBudget | None = None) -> list[Document]:
    """Read every supported member of a zip. Titles are the member paths inside the archive."""
    budget = budget or TextBudget()
    docs: list[Document] = []
    with zipfile.ZipFile(path) as archive:
        members = [m for m in sorted(archive.infolist(), key=lambda m: m.filename) if _wanted_zip_member(m)]
        check_zip_budgets(members, name)
        for member in members:
            try:
                data = archive.read(member)
                docs.extend(
                    read_bytes(data, member.filename, path=f"{name}:{member.filename}", budget=budget)
                )
            except TooLarge:
                raise  # the whole source is over budget; do not treat it as one bad member
            except Exception as err:  # noqa: BLE001 - one broken member must not sink the archive
                log.warning("Skipping %s in %s: %s", member.filename, name, err)
    return docs


def check_zip_budgets(members: list[zipfile.ZipInfo], name: str) -> None:
    """
    Refuse an archive before reading it when its readable members are too many or, unpacked,
    too big. `file_size` is the unpacked size from the zip's own table, so this costs nothing.
    """
    if len(members) > MAX_ZIP_MEMBERS:
        raise TooLarge(
            f"too large: {name} holds {len(members):,} readable files; the limit is {MAX_ZIP_MEMBERS:,}"
        )
    total = sum(m.file_size for m in members)
    if total > MAX_ZIP_TOTAL_BYTES:
        raise TooLarge(
            f"too large: the readable files in {name} unpack to {total:,} bytes; "
            f"the limit is {MAX_ZIP_TOTAL_BYTES:,}"
        )


def _wanted_zip_member(member: zipfile.ZipInfo) -> bool:
    if member.is_dir() or member.file_size > MAX_FILE_BYTES:
        return False
    parts = PurePosixPath(member.filename).parts
    if any(p in IGNORED_DIRS or p.startswith(".") or p == "__MACOSX" for p in parts[:-1]):
        return False
    if parts[-1].startswith(".") and not is_supported_name(parts[-1]):
        return False
    return is_supported_name(parts[-1])


# --------------------------------------------------------- rich formats


def _read_pdf(data: bytes, budget: TextBudget, where: str) -> str:
    from pypdf import PdfReader

    pages: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")  # many "encrypted" PDFs just have an empty owner password
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            budget.add(len(text), where)  # page by page, so a bloated PDF stops early
            pages.append(text)
    except TooLarge:
        raise
    except Exception as err:  # pypdf raises many different types
        raise ReadError(f"could not read PDF: {err}") from err
    return "\n\n".join(p for p in pages if p)


def _read_docx(data: bytes) -> str:
    import docx

    _check_docx_size(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as err:
        raise ReadError(f"could not read .docx: {err}") from err
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # Turn Word headings into markdown headings so the chunker can split on them.
        level = _heading_level(paragraph.style.name if paragraph.style is not None else "")
        blocks.append(f"{'#' * level} {text}" if level else text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _check_docx_size(data: bytes) -> None:
    """A .docx is a zip; python-docx parses word/document.xml in one go, so look at its size first."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            for member in package.infolist():
                if member.filename == "word/document.xml" and member.file_size > MAX_DOCX_XML_BYTES:
                    raise TooLarge(
                        f"too large: the document body is {member.file_size:,} bytes of XML; "
                        f"the limit is {MAX_DOCX_XML_BYTES:,}"
                    )
    except zipfile.BadZipFile as err:
        raise ReadError(f"could not read .docx: {err}") from err


def _heading_level(style_name: str) -> int:
    if style_name.lower() == "title":
        return 1
    if style_name.lower().startswith("heading"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        return min(int(digits), 6) if digits else 2
    return 0


def _read_epub(data: bytes, budget: TextBudget, where: str) -> str:
    """An epub is a zip of XHTML chapters; read them in reading (spine) order."""
    texts: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as book:
            for name in _epub_chapter_names(book):
                if book.getinfo(name).file_size > MAX_FILE_BYTES:
                    log.warning(
                        "Skipping chapter %s of %s: bigger than %d bytes", name, where, MAX_FILE_BYTES
                    )
                    continue
                text = html_to_text(_decode(book.read(name), strict=False))
                budget.add(len(text), f"{where}:{name}")  # chapter by chapter, so a bloated book stops early
                texts.append(text)
    except zipfile.BadZipFile as err:
        raise ReadError(f"could not read .epub: {err}") from err
    return "\n\n".join(t for t in texts if t.strip())


def _epub_chapter_names(book: zipfile.ZipFile) -> list[str]:
    """Chapter file names in spine order; falls back to every html file, sorted."""
    import xml.etree.ElementTree as ET

    names = book.namelist()
    opf_name = next((n for n in names if n.lower().endswith(".opf")), None)
    ordered: list[str] = []
    if opf_name:
        try:
            root = ET.fromstring(book.read(opf_name))
            base = PurePosixPath(opf_name).parent
            hrefs = {
                item.get("id"): str(base / item.get("href", ""))
                for item in root.iter()
                if item.tag.endswith("item") and item.get("id")
            }
            for ref in root.iter():
                if ref.tag.endswith("itemref") and ref.get("idref") in hrefs:
                    ordered.append(hrefs[ref.get("idref")].lstrip("./"))
        except ET.ParseError:
            ordered = []
    known = set(names)
    ordered = [n for n in ordered if n in known]
    if ordered:
        return ordered
    return sorted(n for n in names if _suffix(n) in (".html", ".htm", ".xhtml", ".xml"))


# --------------------------------------------------------------- helpers


def _suffix(name: str) -> str:
    return PurePosixPath(name).suffix.lower()


def _decode(data: bytes, *, strict: bool) -> str:
    """UTF-8 (with or without BOM). Non-strict mode replaces bad bytes instead of failing."""
    return data.decode("utf-8-sig", errors="strict" if strict else "replace")
